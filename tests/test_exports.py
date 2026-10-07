"""导出功能测试：Excel / Markdown。"""
import openpyxl

from core.exports import export_excel, export_markdown
from testcases.models import TestCase


def _make_cases(requirement):
    return [
        TestCase(
            requirement=requirement,
            case_id="TC_001",
            module="登录模块",
            title="登录成功",
            preconditions=["账号已注册"],
            steps=["输入手机号", "点击登录"],
            test_data="手机号=13800138000",
            expected_result="跳转首页",
            test_type="normal",
            priority="P0",
            design_method="scenario",
        )
    ]


class TestExports:
    def test_excel_contains_headers_and_rows(self, requirement):
        content = export_excel(_make_cases(requirement))
        wb = openpyxl.load_workbook(filename=__import__("io").BytesIO(content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert "用例编号" in headers
        assert "测试步骤" in headers
        assert ws.max_row == 2  # 表头 + 1 条数据
        # 步骤数组应转换为人类可读格式
        assert "1. 输入手机号" in str(ws.cell(row=2, column=5).value)

    def test_markdown_contains_case(self, requirement):
        md = export_markdown(requirement, _make_cases(requirement))
        assert "# 测试用例：用户登录" in md
        assert "TC_001" in md
